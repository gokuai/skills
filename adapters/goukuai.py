#!/usr/bin/env python3
"""
够快云库适配器

实现够快云库 API 的标准适配器
"""

import os
import time
import hmac
import hashlib
import base64
import json
import urllib.request
import urllib.error
from typing import List, Dict, Any, Optional
from urllib.parse import quote

from .base import CloudAdapter, FileInfo, FolderInfo, ShareLink, Permission


class GoukuaiAdapter(CloudAdapter):
    """
    够快云库适配器
    
    实现够开云库 YK3 API 的标准接口
    """
    
    def __init__(self, credentials: Dict[str, Any]):
        """
        初始化适配器
        
        Args:
            credentials: 认证信息
                - client_id: 企业 client_id
                - client_secret: 企业 client_secret
                - org_client_id: 库 client_id（可选）
                - org_client_secret: 库 client_secret（可选）
                - api_host: API 域名（默认：yk3.gokuai.com）
        """
        super().__init__(credentials)
        
        self.client_id = credentials.get("client_id")
        self.client_secret = credentials.get("client_secret")
        self.org_client_id = credentials.get("org_client_id")
        self.org_client_secret = credentials.get("org_client_secret")
        self.api_host = credentials.get("api_host", "yk3.gokuai.com")
        
        # 使用库级认证还是企业级认证
        self.use_org_auth = bool(self.org_client_id and self.org_client_secret)
    
    def _calculate_sign(self, params: Dict[str, Any], secret: str = None) -> str:
        """
        计算 API 签名
        
        Args:
            params: 请求参数
            secret: 密钥（默认使用 client_secret 或 org_client_secret）
        
        Returns:
            str: Base64 编码的签名
        """
        # 选择密钥
        if secret is None:
            secret = self.org_client_secret if self.use_org_auth else self.client_secret
        
        # 1. 按字母顺序排序参数（排除 sign 和文件内容）
        sorted_keys = sorted([k for k in params.keys() if k not in ['sign', 'file']])
        
        # 2. 拼接参数值
        values = []
        for key in sorted_keys:
            value = params[key]
            if isinstance(value, (list, dict)):
                value = json.dumps(value, separators=(',', ':'))
            values.append(str(value))
        
        sign_string = '\n'.join(values)
        
        # 3. HMAC-SHA1 加密
        hmac_code = hmac.new(
            secret.encode('utf-8'),
            sign_string.encode('utf-8'),
            hashlib.sha1
        ).digest()
        
        # 4. Base64 编码
        sign = base64.b64encode(hmac_code).decode('utf-8')
        
        return sign
    
    def _request(self, endpoint: str, params: Dict[str, Any], 
                 method: str = "POST", use_org_secret: bool = True) -> Dict:
        """
        发送 API 请求
        
        Args:
            endpoint: API 端点
            params: 请求参数
            method: 请求方法
            use_org_secret: 是否使用 org_client_secret 签名
        
        Returns:
            Dict: 响应结果
        """
        url = f"https://{self.api_host}{endpoint}"
        
        # 添加时间戳
        params['dateline'] = int(time.time())
        
        # 选择密钥
        secret = self.org_client_secret if use_org_secret else self.client_secret
        
        # 计算签名
        params['sign'] = self._calculate_sign(params, secret)
        
        # 构建请求
        headers = {"Content-Type": "application/x-www-form-urlencoded; charset=UTF-8"}
        
        # URL 编码参数
        encoded_params = []
        for key, value in params.items():
            if isinstance(value, (list, dict)):
                value = json.dumps(value, separators=(',', ':'))
            encoded_value = quote(str(value), safe='')
            encoded_params.append(f"{key}={encoded_value}")
        
        data = '&'.join(encoded_params).encode('utf-8')
        
        # 发送请求
        req = urllib.request.Request(url, data=data, headers=headers, method=method)
        
        try:
            with urllib.request.urlopen(req, timeout=30) as response:
                result = json.loads(response.read().decode('utf-8'))
                return result
        except urllib.error.HTTPError as e:
            error_body = e.read().decode('utf-8') if e.fp else ""
            raise Exception(f"API 请求失败：{e.code} - {e.reason}\n{error_body}")
        except urllib.error.URLError as e:
            raise Exception(f"网络错误：{e.reason}")
    
    def authenticate(self) -> bool:
        """验证认证信息"""
        try:
            # 尝试获取库信息来验证
            if self.use_org_auth:
                self.get_library_info()
            else:
                # 企业级认证，尝试获取成员列表
                self._request("/m-open/1/ent/get_members", {
                    "client_id": self.client_id,
                    "start": 0,
                    "size": 1
                }, use_org_secret=False)
            return True
        except Exception as e:
            print(f"认证失败：{e}")
            return False
    
    def get_library_info(self) -> Dict[str, Any]:
        """获取文件库信息"""
        if not self.use_org_auth:
            raise Exception("需要配置 org_client_id 和 org_client_secret")
        
        # 通过统计接口获取库信息
        stat = self.get_library_stat()
        
        return {
            "org_id": stat.get("org_id"),
            "org_name": stat.get("org_name"),
            "mount_id": stat.get("mount_id"),
            "size_org_total": stat.get("capacity"),
            "size_org_use": stat.get("size"),
            "count_file": stat.get("count_file")
        }
    
    def list_files(self, fullpath: str = "", tag: Optional[str] = None,
                   start: int = 0, size: int = 100) -> List[FileInfo]:
        """列出文件夹内容"""
        params = {
            "org_client_id": self.org_client_id,
            "fullpath": fullpath,
            "start": start,
            "size": size
        }
        
        if tag:
            params['tag'] = tag
        
        result = self._request("/m-open/1/file/ls", params)
        
        files = []
        for item in result.get("list", []):
            file_info = FileInfo(
                hash=item.get("hash", ""),
                dir=bool(item.get("dir", 0)),
                fullpath=item.get("fullpath", ""),
                filename=item.get("filename", ""),
                filehash=item.get("filehash"),
                filesize=item.get("filesize", 0),
                create_member_name=item.get("create_member_name"),
                create_dateline=item.get("create_dateline"),
                last_member_name=item.get("last_member_name"),
                last_dateline=item.get("last_dateline"),
                tag=item.get("property", {}).get("tag"),
                permission=item.get("property", {}).get("permission")
            )
            files.append(file_info)
        
        return files
    
    def search_files(self, keywords: str, path: str = "",
                     scope: List[str] = None, start: int = 0, size: int = 100) -> List[FileInfo]:
        """搜索文件"""
        if scope is None:
            scope = ["filename", "tag"]
        
        params = {
            "org_client_id": self.org_client_id,
            "keywords": keywords,
            "path": path,
            "scope": json.dumps(scope),
            "start": start,
            "size": size
        }
        
        result = self._request("/m-open/1/file/search", params)
        
        files = []
        for item in result.get("list", []):
            file_info = FileInfo(
                hash=item.get("hash", ""),
                dir=bool(item.get("dir", 0)),
                fullpath=item.get("fullpath", ""),
                filename=item.get("filename", ""),
                filehash=item.get("filehash"),
                filesize=item.get("filesize", 0),
                create_member_name=item.get("create_member_name"),
                create_dateline=item.get("create_dateline"),
                last_member_name=item.get("last_member_name"),
                last_dateline=item.get("last_dateline")
            )
            files.append(file_info)
        
        return files
    
    def download_file(self, file_hash: str = None, fullpath: str = None,
                      open_browser: bool = False) -> bytes:
        """下载文件"""
        if not file_hash and not fullpath:
            raise Exception("需要指定 file_hash 或 fullpath")
        
        params = {
            "org_client_id": self.org_client_id,
            "open": 1 if open_browser else 0
        }
        
        if file_hash:
            params['hash'] = file_hash
        elif fullpath:
            params['fullpath'] = fullpath
        
        # 添加操作人信息
        if self.current_user_id:
            params['op_id'] = self.current_user_id
        elif self.current_user_account:
            params['op_name'] = self.current_user_account
        
        result = self._request("/m-open/1/file/download_url", params)
        
        urls = result.get("urls", [])
        if not urls:
            raise Exception("未获取到下载链接")
        
        # 下载第一个链接
        try:
            with urllib.request.urlopen(urls[0], timeout=30) as response:
                return response.read()
        except Exception as e:
            raise Exception(f"下载失败：{e}")
    
    def upload_file(self, fullpath: str, content: bytes,
                    overwrite: bool = False) -> FileInfo:
        """上传文件"""
        import hashlib
        
        # 计算文件 hash
        filehash = hashlib.sha1(content).hexdigest()
        filesize = len(content)
        
        # 步骤 1: 请求上传
        params = {
            "org_client_id": self.org_client_id,
            "fullpath": fullpath,
            "filehash": filehash,
            "filesize": filesize,
            "overwrite": 1 if overwrite else 0
        }
        
        # 添加操作人信息
        if self.current_user_id:
            params['op_id'] = self.current_user_id
        elif self.current_user_account:
            params['op_name'] = self.current_user_account
        
        result = self._request("/m-open/1/file/create_file", params)
        
        # 检查是否已秒传
        if result.get("state") == 1:
            return FileInfo(
                hash=result.get("hash"),
                dir=False,
                fullpath=result.get("fullpath"),
                filename=fullpath.split("/")[-1],
                filehash=filehash,
                filesize=filesize
            )
        
        # 步骤 2: 分块上传（简化实现，直接上传）
        # 实际应该按照 /yk3/upload.md 实现分块上传
        upload_server = result.get("server")
        
        # 获取上传密钥
        upload_params = {
            "org_client_id": self.org_client_id,
            "fullpath": "/".join(fullpath.split("/")[:-1]) or "",
            "timeout": 300,
            "rand": hashlib.md5(str(time.time()).encode()).hexdigest()[:16]
        }
        
        upload_result = self._request("/m-open/1/file/upload_servers", upload_params)
        
        upload_key = upload_result.get("key")
        upload_servers = upload_result.get("m-upload", [])
        
        if not upload_servers:
            raise Exception("未获取到上传服务器")
        
        # 上传文件
        import urllib.parse
        boundary = '----WebKitFormBoundary7MA4YWxkTrZu0gW'
        
        body = []
        body.append(f'--{boundary}'.encode())
        body.append(b'Content-Disposition: form-data; name="org_client_id"')
        body.append(b'')
        body.append(self.org_client_id.encode())
        
        body.append(f'--{boundary}'.encode())
        body.append(b'Content-Disposition: form-data; name="key"')
        body.append(b'')
        body.append(upload_key.encode())
        
        body.append(f'--{boundary}'.encode())
        body.append(b'Content-Disposition: form-data; name="path"')
        body.append(b'')
        body.append(b'')
        
        body.append(f'--{boundary}'.encode())
        body.append(b'Content-Disposition: form-data; name="name"')
        body.append(b'')
        body.append(fullpath.split("/")[-1].encode())
        
        body.append(f'--{boundary}'.encode())
        body.append(b'Content-Disposition: form-data; name="overwrite"')
        body.append(b'')
        body.append(b'1' if overwrite else b'0')
        
        body.append(f'--{boundary}'.encode())
        body.append(b'Content-Disposition: form-data; name="filefield"')
        body.append(b'')
        body.append(b'file')
        
        body.append(f'--{boundary}'.encode())
        body.append(f'Content-Disposition: form-data; name="file"; filename="{fullpath.split("/")[-1]}"'.encode())
        body.append(b'Content-Type: application/octet-stream')
        body.append(b'')
        body.append(content)
        
        body.append(f'--{boundary}--'.encode())
        
        body_data = b'\r\n'.join(body)
        
        headers = {
            "Content-Type": f"multipart/form-data; boundary={boundary}",
            "Content-Length": str(len(body_data))
        }
        
        upload_url = upload_servers[0] + "/2/web_upload"
        req = urllib.request.Request(upload_url, data=body_data, headers=headers, method="POST")
        
        try:
            with urllib.request.urlopen(req, timeout=60) as response:
                result = json.loads(response.read().decode('utf-8'))
                return FileInfo(
                    hash=result.get("hash", ""),
                    dir=False,
                    fullpath=result.get("fullpath", ""),
                    filename=result.get("fullpath", "").split("/")[-1],
                    filehash=filehash,
                    filesize=result.get("filesize", 0)
                )
        except Exception as e:
            raise Exception(f"上传失败：{e}")
    
    def get_file_content(self, file_hash: str = None, fullpath: str = None) -> str:
        """获取文件文本内容"""
        # 下载文件
        content = self.download_file(file_hash=file_hash, fullpath=fullpath)
        
        # 根据文件类型提取文本
        filename = fullpath or file_hash
        ext = filename.split(".")[-1].lower() if "." in filename else ""
        
        if ext == "pdf":
            return self._extract_pdf_text(content)
        elif ext in ["docx", "doc"]:
            return self._extract_word_text(content)
        elif ext in ["xlsx", "xls"]:
            return self._extract_excel_text(content)
        elif ext in ["txt", "md", "json", "csv"]:
            return content.decode('utf-8', errors='ignore')
        else:
            # 默认尝试 UTF-8 解码
            return content.decode('utf-8', errors='ignore')
    
    def _extract_pdf_text(self, content: bytes) -> str:
        """提取 PDF 文本"""
        try:
            import pypdf
            from io import BytesIO
            
            reader = pypdf.PdfReader(BytesIO(content))
            text = []
            for page in reader.pages:
                text.append(page.extract_text())
            return '\n'.join(text)
        except ImportError:
            return "[PDF 内容需要安装 pypdf 库]"
        except Exception as e:
            return f"[PDF 提取失败：{e}]"
    
    def _extract_word_text(self, content: bytes) -> str:
        """提取 Word 文本"""
        try:
            from docx import Document
            from io import BytesIO
            
            doc = Document(BytesIO(content))
            text = [para.text for para in doc.paragraphs]
            return '\n'.join(text)
        except ImportError:
            return "[Word 内容需要安装 python-docx 库]"
        except Exception as e:
            return f"[Word 提取失败：{e}]"
    
    def _extract_excel_text(self, content: bytes) -> str:
        """提取 Excel 文本"""
        try:
            import openpyxl
            from io import BytesIO
            
            wb = openpyxl.load_workbook(BytesIO(content), read_only=True)
            text = []
            for sheet in wb.worksheets:
                sheet_text = []
                for row in sheet.iter_rows(values_only=True):
                    row_text = [str(cell) if cell is not None else "" for cell in row]
                    sheet_text.append("\t".join(row_text))
                text.append(f"=== {sheet.title} ===\n" + "\n".join(sheet_text))
            return '\n'.join(text)
        except ImportError:
            return "[Excel 内容需要安装 openpyxl 库]"
        except Exception as e:
            return f"[Excel 提取失败：{e}]"
    
    def get_file_info(self, file_hash: str = None, fullpath: str = None,
                      attribute: bool = False) -> FileInfo:
        """获取文件（夹）信息"""
        if not file_hash and not fullpath:
            raise Exception("需要指定 file_hash 或 fullpath")
        
        params = {
            "org_client_id": self.org_client_id,
            "attribute": "1" if attribute else "0"
        }
        
        if file_hash:
            params['hash'] = file_hash
        elif fullpath:
            params['fullpath'] = fullpath
        
        result = self._request("/m-open/1/file/info", params)
        
        prop = result.get("property", {})
        
        return FileInfo(
            hash=result.get("hash", ""),
            dir=bool(result.get("dir", 0)),
            fullpath=result.get("fullpath", ""),
            filename=result.get("filename", ""),
            filehash=result.get("filehash"),
            filesize=result.get("filesize", 0),
            create_member_name=result.get("create_member_name"),
            create_dateline=result.get("create_dateline"),
            last_member_name=result.get("last_member_name"),
            last_dateline=result.get("last_dateline"),
            tag=prop.get("tag"),
            permission=prop.get("permission")
        )
    
    def copy_file(self, from_fullpath: str, to_fullpath: str,
                  overwrite: bool = False) -> FileInfo:
        """复制文件（夹）"""
        params = {
            "org_client_id": self.org_client_id,
            "from_fullpath": from_fullpath,
            "fullpath": to_fullpath,
            "overwrite": 1 if overwrite else 0
        }
        
        # 添加操作人信息
        if self.current_user_id:
            params['op_id'] = self.current_user_id
        elif self.current_user_account:
            params['op_name'] = self.current_user_account
        
        result = self._request("/m-open/1/file/copy", params)
        
        # 检查是否需要异步处理
        queue_id = result.get("queue_id")
        if queue_id:
            # 异步处理，查询队列状态
            return self._wait_queue_complete(queue_id)
        
        return FileInfo(
            hash=result.get("hash", ""),
            dir=False,
            fullpath=result.get("fullpath", ""),
            filename=to_fullpath.split("/")[-1]
        )
    
    def _wait_queue_complete(self, queue_id: str, timeout: int = 60) -> FileInfo:
        """等待队列任务完成"""
        start_time = time.time()
        
        while time.time() - start_time < timeout:
            result = self._request("/m-open/1/file/queue_status", {
                "org_client_id": self.org_client_id,
                "queue_id": queue_id
            })
            
            status = result.get("status", 0)
            
            if status == 2:  # 完成
                return FileInfo(
                    hash="",
                    dir=False,
                    fullpath=result.get("fullpath", "")
                )
            elif status == 3:  # 出错
                raise Exception("队列任务执行失败")
            
            time.sleep(2)
        
        raise Exception("等待队列任务超时")
    
    def move_file(self, from_fullpath: str, to_fullpath: str) -> bool:
        """移动文件（夹）"""
        params = {
            "org_client_id": self.org_client_id,
            "fullpath": from_fullpath,
            "dest_fullpath": to_fullpath
        }
        
        # 添加操作人信息
        if self.current_user_id:
            params['op_id'] = self.current_user_id
        elif self.current_user_account:
            params['op_name'] = self.current_user_account
        
        self._request("/m-open/1/file/move", params)
        return True
    
    def delete_file(self, fullpath: str = None, tag: str = None,
                    destroy: bool = False) -> bool:
        """删除文件（夹）"""
        if not fullpath and not tag:
            raise Exception("需要指定 fullpath 或 tag")
        
        params = {
            "org_client_id": self.org_client_id,
            "destroy": 1 if destroy else 0
        }
        
        if fullpath:
            params['fullpath'] = fullpath
        elif tag:
            params['tag'] = tag
        
        # 添加操作人信息
        if self.current_user_id:
            params['op_id'] = self.current_user_id
        elif self.current_user_account:
            params['op_name'] = self.current_user_account
        
        self._request("/m-open/1/file/del", params)
        return True
    
    def create_folder(self, fullpath: str) -> FolderInfo:
        """创建文件夹"""
        params = {
            "org_client_id": self.org_client_id,
            "fullpath": fullpath
        }
        
        # 添加操作人信息
        if self.current_user_id:
            params['op_id'] = self.current_user_id
        elif self.current_user_account:
            params['op_name'] = self.current_user_account
        
        result = self._request("/m-open/1/file/create_folder", params)
        
        return FolderInfo(
            hash=result.get("hash", ""),
            fullpath=result.get("fullpath", ""),
            filename=fullpath.split("/")[-1]
        )
    
    def get_permission(self, fullpath: str, member_id: str = None) -> List[Permission]:
        """获取文件权限"""
        params = {
            "org_client_id": self.org_client_id,
            "fullpath": fullpath
        }
        
        if member_id:
            params['member_id'] = member_id
        
        result = self._request("/m-open/1/file/get_permission", params)
        
        if member_id:
            # 返回指定用户的权限
            return [Permission(
                member_id=member_id,
                permissions=result if isinstance(result, list) else []
            )]
        else:
            # 返回所有权限
            permissions = []
            
            # 成员权限
            members = result.get("members", {})
            for mid, perms in members.items():
                permissions.append(Permission(
                    member_id=mid,
                    permissions=perms
                ))
            
            # 部门权限
            groups = result.get("groups", {})
            for gid, perms in groups.items():
                permissions.append(Permission(
                    group_id=gid,
                    permissions=perms
                ))
            
            return permissions
    
    def set_permission(self, fullpath: str, permissions: Dict) -> bool:
        """设置文件权限"""
        params = {
            "org_client_id": self.org_client_id,
            "fullpath": fullpath,
            "permissions": json.dumps(permissions)
        }
        
        self._request("/m-open/2/file/set_permission", params)
        return True
    
    def set_permission_inherit(self, fullpath: str, inherit: bool,
                                keep: bool = False) -> bool:
        """设置权限继承"""
        params = {
            "org_client_id": self.org_client_id,
            "fullpath": fullpath,
            "inherit": 1 if inherit else 0,
            "keep": 1 if keep else 0
        }
        
        self._request("/m-open/2/file/set_permission_inherit", params)
        return True
    
    def create_share_link(self, fullpath: str, deadline: int = None,
                          password: str = None, auth: str = "preview",
                          dir: bool = False) -> ShareLink:
        """创建分享链接"""
        params = {
            "org_client_id": self.org_client_id,
            "fullpath": fullpath,
            "auth": auth,
            "dir": 1 if dir else 0
        }
        
        if deadline:
            params['deadline'] = deadline
        if password:
            params['password'] = password
        
        # 添加操作人信息
        if self.current_user_id:
            params['op_id'] = self.current_user_id
        elif self.current_user_account:
            params['op_name'] = self.current_user_account
        
        result = self._request("/m-open/1/file/link", params)
        
        return ShareLink(
            link=result.get("link", ""),
            code=result.get("code", ""),
            deadline=deadline,
            password=password,
            auth=auth
        )
    
    def close_share_link(self, code: str = None, fullpath: str = None) -> bool:
        """关闭分享链接"""
        if not code and not fullpath:
            raise Exception("需要指定 code 或 fullpath")
        
        params = {
            "org_client_id": self.org_client_id
        }
        
        if code:
            params['code'] = code
        elif fullpath:
            params['fullpath'] = fullpath
        
        # 添加操作人信息
        if self.current_user_id:
            params['op_id'] = self.current_user_id
        elif self.current_user_account:
            params['op_name'] = self.current_user_account
        
        self._request("/m-open/1/file/link_close", params)
        return True
    
    def get_cedit_url(self, file_hash: str = None, fullpath: str = None,
                      readonly: bool = False, timeout: int = 3600,
                      op_id: str = None) -> str:
        """获取协同编辑链接"""
        if not file_hash and not fullpath:
            raise Exception("需要指定 file_hash 或 fullpath")
        
        params = {
            "org_client_id": self.org_client_id,
            "readonly": 1 if readonly else 0,
            "timeout": timeout
        }
        
        if file_hash:
            params['hash'] = file_hash
        elif fullpath:
            params['fullpath'] = fullpath
        
        # 操作人必须指定
        params['op_id'] = op_id or self.current_user_id or ""
        
        result = self._request("/m-open/1/file/cedit_url", params)
        return result.get("url", "")
    
    def get_history(self, fullpath: str, start: int = 0, size: int = 20) -> List[Dict]:
        """获取文件历史版本"""
        params = {
            "org_client_id": self.org_client_id,
            "fullpath": fullpath,
            "start": start,
            "size": size
        }
        
        result = self._request("/m-open/1/file/history", params)
        return result.get("list", [])
    
    def recover_version(self, hid: str) -> bool:
        """恢复历史版本"""
        # 够快 API 未直接提供版本恢复接口
        # 需要通过下载历史版本并覆盖实现
        raise NotImplementedError("版本恢复功能需要额外实现")
    
    def add_tag(self, fullpath: str, tag: str) -> bool:
        """添加标签"""
        params = {
            "org_client_id": self.org_client_id,
            "fullpath": fullpath,
            "tag": tag
        }
        
        # 添加操作人信息
        if self.current_user_id:
            params['op_id'] = self.current_user_id
        elif self.current_user_account:
            params['op_name'] = self.current_user_account
        
        self._request("/m-open/1/file/add_tag", params)
        return True
    
    def del_tag(self, fullpath: str, tag: str) -> bool:
        """删除标签"""
        params = {
            "org_client_id": self.org_client_id,
            "fullpath": fullpath,
            "tag": tag
        }
        
        # 添加操作人信息
        if self.current_user_id:
            params['op_id'] = self.current_user_id
        elif self.current_user_account:
            params['op_name'] = self.current_user_account
        
        self._request("/m-open/1/file/del_tag", params)
        return True
    
    def get_library_stat(self) -> Dict[str, Any]:
        """获取统计信息"""
        result = self._request("/m-open/1/file/stat", {
            "org_client_id": self.org_client_id
        })
        return result
    
    # ========== AI 增强方法 ==========
    
    def summarize_file(self, fullpath: str) -> str:
        """摘要文件内容"""
        # 获取文件内容
        content = self.get_file_content(fullpath=fullpath)
        
        # 调用 AI 摘要（需要配置 AI 服务）
        # 这里使用简单的摘要逻辑
        lines = content.split('\n')
        
        # 取前 10 行作为摘要
        summary_lines = []
        for line in lines[:10]:
            if line.strip():
                summary_lines.append(line.strip())
        
        summary = '\n'.join(summary_lines)
        
        if len(lines) > 10:
            summary += f"\n...（共{len(lines)}行）"
        
        return summary
    
    def ask_file(self, fullpath: str, question: str) -> str:
        """基于文件内容问答"""
        # 获取文件内容
        content = self.get_file_content(fullpath=fullpath)
        
        # 调用 AI 问答（需要配置 AI 服务）
        # 这里返回提示信息
        return f"[问答功能需要配置 AI 服务]\n\n文件内容长度：{len(content)} 字符\n问题：{question}"
